from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

title = doc.add_heading('拓扑排序应用软件 中期汇报演示流程', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('')

# 一、开头介绍
doc.add_heading('一、开头介绍', level=1)
doc.add_paragraph('1. 打开命令行，输入 java -cp out ui.MainFrame 启动软件')
doc.add_paragraph('2. 展示主界面：左侧输入区、中间关系图、右侧结果列表、底部状态栏')
doc.add_paragraph('旁白：这是我们的拓扑排序应用软件，输入课程先修关系，自动绘制关系图并枚举所有可能的拓扑排序结果')
doc.add_paragraph('截图：主界面全景截图')

# 二、载入图1数据
doc.add_heading('二、载入图1数据', level=1)
doc.add_paragraph('1. 点【载入文件】按钮')
doc.add_paragraph('2. 选择 data/figure1.txt')
doc.add_paragraph('3. 左侧文本区显示所有 <a,b> 格式的边')
doc.add_paragraph('4. 右侧表格自动同步显示所有边')
doc.add_paragraph('5. 中间关系图自动环形布局，所有节点蓝色')
doc.add_paragraph('旁白：我们先载入任务书图1的课程先修关系数据，支持文本和表格两种编辑方式，关系图自动环形布局')
doc.add_paragraph('截图：载入后完整界面（左文本+右表格+中关系图）')

# 三、计算拓扑排序
doc.add_heading('三、计算拓扑排序', level=1)
doc.add_paragraph('1. 点【计算】按钮')
doc.add_paragraph('2. 状态栏显示：节点数、边数、无环、序列数、耗时')
doc.add_paragraph('3. 右侧结果列表分页显示所有拓扑排序结果')
doc.add_paragraph('4. 切换分页（上一页/下一页）')
doc.add_paragraph('旁白：点击计算，算法自动枚举所有可能的拓扑排序，默认最多显示1000条，支持分页查看')
doc.add_paragraph('截图：计算完成后的结果列表+状态栏')

# 四、高亮选中序列
doc.add_heading('四、高亮选中序列', level=1)
doc.add_paragraph('1. 单击第1条结果 → 全图节点变绿色，每个节点右上角标序号徽章（1,2,3,4...）')
doc.add_paragraph('2. 单击第2条结果 → 全图变橙色')
doc.add_paragraph('3. 单击第3条结果 → 全图变紫色')
doc.add_paragraph('4. 双击第1条结果 → 复制到剪贴板')
doc.add_paragraph('旁白：单击任意结果，关系图会高亮显示这条拓扑序列，不同结果用不同颜色区分，节点上的序号表示该节点在序列中的位置，双击可复制')
doc.add_paragraph('截图1：全图绿色高亮+节点右上角序号徽章')
doc.add_paragraph('截图2：全图变成橙色（点第2条结果）')

# 五、环检测演示
doc.add_heading('五、环检测演示', level=1)
doc.add_paragraph('1. 左侧文本区清空，输入3条有环的边：')
doc.add_paragraph('   <A,B>')
doc.add_paragraph('   <B,C>')
doc.add_paragraph('   <C,A>')
doc.add_paragraph('2. 点【计算】')
doc.add_paragraph('3. 弹窗提示「图中存在环，无法进行拓扑排序」')
doc.add_paragraph('4. 关系图上环上的节点和边全部变红')
doc.add_paragraph('5. 状态栏显示「含环」')
doc.add_paragraph('旁白：如果输入有环的图，系统会自动检测并红色高亮显示环路径，无法进行拓扑排序')
doc.add_paragraph('截图：环上节点和边全部变红的界面')

# 六、导出功能演示
doc.add_heading('六、导出功能演示', level=1)
doc.add_paragraph('1. 切回 data/figure1.txt，点计算')
doc.add_paragraph('2. 点【导出结果】→ 保存为 txt 文件')
doc.add_paragraph('3. 点【导出图片】→ 保存关系图为 PNG')
doc.add_paragraph('旁白：支持导出排序结果为txt文件，也支持导出关系图为PNG图片')
doc.add_paragraph('截图1：导出的txt文件内容')
doc.add_paragraph('截图2：导出的PNG关系图')

# 七、异常输入处理
doc.add_heading('七、异常输入处理', level=1)
doc.add_paragraph('1. 左侧输入 <A,>（不完整的边）')
doc.add_paragraph('2. 点计算 → 弹窗提示「第1行：...」')
doc.add_paragraph('旁白：系统有完善的异常处理，非法输入会给出明确的中文错误提示和行号定位')
doc.add_paragraph('截图：错误提示弹窗')

# 八、录制备注
doc.add_heading('八、录制备注', level=1)
doc.add_paragraph('1. 录屏分辨率 1080p')
doc.add_paragraph('2. 鼠标移动慢一点，点按钮时停1秒让观众看清')
doc.add_paragraph('3. 旁白声音清晰，不要快')
doc.add_paragraph('4. 总时长控制在5分钟以内')
doc.add_paragraph('5. 结尾停在主界面3秒')

# 九、E的任务
doc.add_heading('九、E的任务清单', level=1)
doc.add_paragraph('1. 按上面演示流程录屏+配音')
doc.add_paragraph('2. 剪辑：开头加标题「拓扑排序应用软件 中期汇报」')
doc.add_paragraph('3. 压缩：控制在100M以内')
doc.add_paragraph('4. 命名：groupXX.mp4（按组号）')
doc.add_paragraph('5. 提交到mystu对应老师的入口（1-21组廖海泳/22-42组陈银冬）')

doc.save(r'E:\tp\中期汇报演示流程.docx')
print('done')
