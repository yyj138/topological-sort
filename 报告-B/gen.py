from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 标题
title = doc.add_heading('3.2 GUI与控制模块详细设计（易雨杰 2024611209）', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 3.2.1 类设计
doc.add_heading('3.2.1 类设计', level=1)
doc.add_paragraph('GUI与控制模块采用MVC架构，共包含以下核心类：')

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = '类名'
hdr[1].text = '所在包'
hdr[2].text = '职责说明'

classes = [
    ('MainFrame', 'ui', '主窗口，三栏布局，持有所有UI组件引用'),
    ('MainController', 'ui', '主控制器，接收UI事件，调度后台线程，对接A/C/D各模块接口'),
    ('InputPanel', 'ui', '左栏输入面板，文本编辑区，文件导入按钮'),
    ('ResultPanel', 'ui', '右栏结果面板，JList分页显示，选中联动事件'),
    ('StatusBar', 'ui', '底部状态栏，显示节点/边/结果统计和状态提示'),
    ('UIStyle', 'util', '统一UI样式工具类，配色、字体、按钮样式'),
    ('ExceptionHandler', 'util', '统一异常弹窗，错误提示'),
]
for c in classes:
    row = table.add_row().cells
    row[0].text = c[0]
    row[1].text = c[1]
    row[2].text = c[2]

doc.add_heading('核心方法说明', level=2)
doc.add_paragraph('MainController 核心方法：')
doc.add_paragraph('- compute()：接收"计算"按钮事件，启动后台线程')
doc.add_paragraph('- applyEnumerationResult()：计算完成后更新界面，把结果传给图和列表')
doc.add_paragraph('- onSequenceSelected()：处理结果选中事件，调用图高亮对应拓扑序')
doc.add_paragraph('- cancelCompute()：中断当前计算线程')
doc.add_paragraph('- statusText()：更新状态栏提示')

# 3.2.2 核心流程
doc.add_heading('3.2.2 核心流程', level=1)
doc.add_paragraph('完整计算与显示流程：')
doc.add_paragraph('1. 用户在InputPanel输入或导入关系数据，点击"计算"按钮')
doc.add_paragraph('2. MainController校验输入非空，创建SwingWorker后台任务')
doc.add_paragraph('3. 后台线程：调用D的DataParser解析文本 → 构建Graph → 调用A的AllTopoSorts枚举拓扑序')
doc.add_paragraph('4. 计算过程中主线程不阻塞，界面可操作，可随时点"取消"中断')
doc.add_paragraph('5. 计算完成回到EDT线程：GraphPanel显示图，ResultPanel显示结果列表，状态栏更新统计')
doc.add_paragraph('6. 用户点击结果列表某行 → ResultPanel回调 → MainController调用GraphPanel.setSelectedOrder高亮该序列')

doc.add_paragraph('结果选中联动流程：')
doc.add_paragraph('1. 用户点击JList某一项，触发ListSelectionListener')
doc.add_paragraph('2. ResultPanel计算全局索引（当前页×每页20 + 行内索引）')
doc.add_paragraph('3. 通过SelectionListener把该序列传给MainController')
doc.add_paragraph('4. MainController把序列和选中序号传给GraphPanel.setSelectedOrder')
doc.add_paragraph('5. GraphPanel按序号选择颜色，高亮对应节点和边')

# 3.2.3 交互设计
doc.add_heading('3.2.3 交互设计', level=1)
doc.add_paragraph('界面采用三栏式布局：')
doc.add_paragraph('- 顶部：JToolBar工具栏，放置常用操作按钮，统一蓝色主题')
doc.add_paragraph('- 中部：JSplitPane水平分割为三栏，各栏宽度可拖拽调整')
doc.add_paragraph('- 底部：StatusBar状态栏')

doc.add_paragraph('交互设计原则：')
doc.add_paragraph('- 计算在后台线程执行，界面不卡顿')
doc.add_paragraph('- 所有错误通过统一弹窗提示，带行号定位')
doc.add_paragraph('- 结果分页显示，避免上百条结果一次性加载')
doc.add_paragraph('- 按钮和面板颜色、字体统一，通过UIStyle配置')

# 个人感想
doc.add_heading('个人感想与收获', level=1)
doc.add_paragraph('作为小组的项目统筹和任务B的负责人，在这次高级算法课程项目中我收获很多。')
doc.add_paragraph('从技术层面来说，这是我第一次完整经历一个多人协作的Java项目开发流程。最开始我们五个人先一起开了几次会，确定了整体架构和接口契约，把功能拆成A到E五个模块，每个人负责一块。我负责的是GUI主框架和交互控制，一开始我以为写Swing界面很简单，就是拖一拖组件，但真正做起来才发现，要把不同人写的模块拼到一起、接口对齐、不打架，才是最花时间的。比如最开始和C对接可视化模块的时候，他写的GraphPanel和我写的MainController之间经常对不上，他改了接口我这边没更新，我这边加了功能他那边不知道，来回沟通了好几次才对齐。')
doc.add_paragraph('和其他组员的对接让我学会了怎么写清晰的接口文档，怎么提前约定好方法名、参数、返回值，而不是写完代码再凑。比如D负责导出，我们之前约定了导出的时候要传"结果是否完整"和"停止原因"两个参数，这样导出的文件才符合任务书要求，不会把截断的结果标成全部。这个过程让我理解了为什么大项目里接口设计这么重要——先把接口定死，大家各自开发，最后再拼，比边写边改效率高太多了。')
doc.add_paragraph('项目中间也遇到了不少问题，比如最开始合并代码的时候冲突一大堆，尤其是GraphPanel这个文件，C改一点我改一点，每次合并都要解决冲突。还有UI乱码的问题，一开始想用特殊符号做按钮图标，结果Swing在Windows下显示成方框，折腾了半天才发现是字体不支持，最后改成中文"重置"才解决。还有放大查看窗口不跟随布局的bug，找了半天才发现是弹新窗口的时候忘了把当前的布局模式传过去。')
doc.add_paragraph('团队合作方面，这次五个人分工明确，我们前后开了四次线上会议。第一次是项目启动会，大家一起把任务拆成五个模块，确认了每个人的分工，搭好GitHub仓库和分支；第二次会议是接口评审，把所有模块之间的方法签名、参数、返回值都定下来，冻结接口，大家开始各自开发；第三次会议是联调会，把各自写好的代码拼到一起，对接了A的算法、C的可视化、D的IO，发现了一堆接口对不上的问题，比如导出的时候缺参数、弹窗不跟随布局，现场调了一下午；第四次会议是中期准备会，把所有代码合并到dev-b分支，录中期演示视频，确认了接下来一周的测试和bug修复安排。')
doc.add_paragraph('一开始用Git分支开发，经常有人推了新代码其他人不知道，导致版本对不上。后来我们约定每次会议都先拉取最新代码，有改动及时在群里说，对接的时候两边先对齐参数再写代码，慢慢就顺畅了。我作为统筹，还要协调大家的进度，提醒谁该做什么，对接的时候两边对齐参数，虽然有点琐碎，但是也锻炼了我组织和沟通的能力。')
doc.add_paragraph('整体来说，这次项目不仅让我把课堂上学到的拓扑排序、图论知识用到了实际项目里，更重要的是学会了怎么和别人一起协作写代码、怎么设计模块接口、怎么处理版本冲突，这些都是课堂上学不到的工程经验。虽然过程中改了很多bug，合并了很多次代码，但是最后看到整个软件能完整跑起来，从输入关系到枚举所有拓扑序到可视化高亮，那种成就感还是很强的。')

doc.save(r'E:\tp\报告-B\易雨杰-3.2-GUI与控制模块.docx')
print('done')
