from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
title = doc.add_heading('项目分工与工作量分配', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('本项目共5人参与，按照"接口冻结、分支开发、联调合并"的原则分工：初期共同制定各模块接口契约，之后各自在独立分支完成开发，中期进行代码集成和测试，后期统一整理文档和交付材料。')

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = '成员'
hdr[1].text = '负责模块'
hdr[2].text = '主要工作内容'

members = [
    ('骆深敏\n2024611026', 'A 算法模块',
     '设计Graph/Vertex图数据结构，实现Kahn单条拓扑排序、DFS回溯全拓扑枚举（上限1000条）、环检测与环路径定位；牵头制定模块间接口契约，完成算法验证和详细设计文档'),
    ('易雨杰\n2024611209', 'B GUI与控制',
     '搭建Swing三栏主窗口和工具栏、状态栏，实现MainController调度和各模块对接；完成结果分页列表、UI统一美化；负责合并各分支代码、协调联调和项目文档整理'),
    ('戴燕岚\n2024611180', 'C 可视化模块',
     '实现GraphPanel节点边绘制、分层/环形双布局；支持缩放平移、节点拖动、拓扑序五色高亮和序号标注；完成右下角悬浮缩放按钮、放大查看弹窗和PNG导出'),
    ('黄佳慧\n2024611022', 'D IO模块',
     '实现输入数据解析和格式校验、错误行号定位；完成TXT/CSV结果导出；整理新版测试数据curriculum.txt，编写异常边界场景清单和测试输出文档'),
    ('吴丽梅\n2024611020', 'E 测试与交付',
     '完成功能测试用例和测试报告，编写readme.txt运行说明；整理四次会议记录，录制中期汇报演示视频；编写用户手册，负责最终工程打包交付'),
]
for m in members:
    row = table.add_row().cells
    row[0].text = m[0]
    row[1].text = m[1]
    row[2].text = m[2]

doc.add_paragraph('开发过程中，MainController作为中间层对接各模块，接口变更需全员同步，代码合并冲突由B协调处理，跨模块对接记录统一存放在docs/协作目录。')

doc.save(r'E:\tp\报告-B\项目分工与工作量分配.docx')
print('done')
