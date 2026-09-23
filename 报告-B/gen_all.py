from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 1. 生成B负责部分-需求分析与GUI设计.docx
doc1 = Document()
title = doc1.add_heading('B负责部分：需求分析与GUI设计', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc1.add_paragraph('易雨杰（2024611209） 任务B')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc1.add_heading('一、需求分析', level=1)
doc1.add_heading('1.1 功能需求', level=2)
table = doc1.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = '需求项'
hdr[1].text = '描述'
hdr[2].text = '优先级'
reqs = [
    ('关系输入', '支持文本框直接粘贴、TXT文件导入、表格编辑三种方式', '高'),
    ('自动建图', '解析<a,b>格式关系，自动去重、识别自环、错误行号提示', '高'),
    ('拓扑排序枚举', '集成Kahn算法+DFS回溯全枚举，默认上限1000条', '高'),
    ('环检测', '自动检测有向环并红色高亮显示环路径', '高'),
    ('可视化展示', '分层/环形双布局，支持缩放、平移、节点拖动', '高'),
    ('结果联动', '点击结果列表，图上同步高亮对应拓扑序，标注顺序序号', '高'),
    ('导出功能', '支持TXT/CSV结果导出、PNG关系图导出', '中'),
    ('稳定保护', '30秒超时、计算过程可取消，大图运算不卡死', '中'),
]
for r in reqs:
    row = table.add_row().cells
    row[0].text = r[0]
    row[1].text = r[1]
    row[2].text = r[2]

doc1.add_heading('1.2 非功能需求', level=2)
doc1.add_paragraph('1. 跨平台：Java Swing实现，Windows/Linux/macOS均可运行')
doc1.add_paragraph('2. 零依赖：纯JDK标准库，无需第三方包')
doc1.add_paragraph('3. 易用性：界面简洁，新用户5分钟可上手')
doc1.add_paragraph('4. 性能：100节点内图秒级响应，1000条结果不卡顿')

doc1.add_heading('二、GUI设计', level=1)
doc1.add_heading('2.1 整体布局', level=2)
doc1.add_paragraph('采用三栏式Swing布局：顶部工具栏、左栏输入、中栏关系图、右栏结果列表、底部状态栏。')

doc1.add_heading('2.2 交互流程', level=2)
doc1.add_paragraph('1. 导入/输入关系数据 → 2. 点击计算后台解析枚举 → 3. 图和结果列表显示 → 4. 点击结果图上联动高亮 → 5. 缩放/平移/切换布局查看 → 6. 导出结果或图片')

doc1.add_heading('2.3 B的实现任务', level=2)
table2 = doc1.add_table(rows=1, cols=3)
table2.style = 'Table Grid'
h = table2.rows[0].cells
h[0].text = '任务编号'
h[1].text = '内容'
h[2].text = '状态'
tasks = [
    ('T-B1', '搭建Swing主窗口框架，三栏布局', '已完成'),
    ('T-B2', '实现MainController主控制器，对接A/C/D各模块', '已完成'),
    ('T-B3', '实现ResultPanel结果面板，分页显示、选中联动', '已完成'),
    ('T-B4', '实现工具栏、导出按钮、取消计算功能', '已完成'),
    ('T-B5', '统一异常弹窗，状态栏状态提示', '已完成'),
    ('T-B6', 'UI美化：UIStyle统一样式，修复乱码', '已完成'),
    ('T-B7', '各模块接口联调，合并到dev-b分支', '已完成'),
]
for t in tasks:
    row = table2.add_row().cells
    row[0].text = t[0]
    row[1].text = t[1]
    row[2].text = t[2]

doc1.save(r'E:\tp\报告-B\B负责部分-需求分析与GUI设计.docx')

# 2. 生成个人感想-易雨杰.docx
doc2 = Document()
title2 = doc2.add_heading('个人感想 —— 易雨杰（2024611209）', 0)
title2.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc2.add_paragraph('作为小组的项目统筹和任务B的负责人，在这次高级算法课程项目中我收获很多。')
doc2.add_paragraph('从技术层面来说，这是我第一次完整经历一个多人协作的Java项目开发流程。最开始我们五个人先一起开了几次会，确定了整体架构和接口契约，把功能拆成A到E五个模块，每个人负责一块。我负责的是GUI主框架和交互控制，一开始我以为写Swing界面很简单，就是拖一拖组件，但真正做起来才发现，要把不同人写的模块拼到一起、接口对齐、不打架，才是最花时间的。比如最开始和C对接可视化模块的时候，他写的GraphPanel和我写的MainController之间经常对不上，他改了接口我这边没更新，我这边加了功能他那边不知道，来回沟通了好几次才对齐。')
doc2.add_paragraph('和其他组员的对接让我学会了怎么写清晰的接口文档，怎么提前约定好方法名、参数、返回值，而不是写完代码再凑。比如D负责导出，我们之前约定了导出的时候要传"结果是否完整"和"停止原因"两个参数，这样导出的文件才符合任务书要求，不会把截断的结果标成全部。这个过程让我理解了为什么大项目里接口设计这么重要——先把接口定死，大家各自开发，最后再拼，比边写边改效率高太多了。')
doc2.add_paragraph('项目中间也遇到了不少问题，比如最开始合并代码的时候冲突一大堆，尤其是GraphPanel这个文件，C改一点我改一点，每次合并都要解决冲突。还有UI乱码的问题，一开始想用特殊符号做按钮图标，结果Swing在Windows下显示成方框，折腾了半天才发现是字体不支持，最后改成中文"重置"才解决。还有放大查看窗口不跟随布局的bug，找了半天才发现是弹新窗口的时候忘了把当前的布局模式传过去。')
doc2.add_paragraph('团队合作方面，这次五个人分工明确，我们前后开了四次线上会议。第一次是项目启动会，大家一起把任务拆成五个模块，确认了每个人的分工，搭好GitHub仓库和分支；第二次会议是接口评审，把所有模块之间的方法签名、参数、返回值都定下来，冻结接口，大家开始各自开发；第三次会议是联调会，把各自写好的代码拼到一起，对接了A的算法、C的可视化、D的IO，发现了一堆接口对不上的问题，比如导出的时候缺参数、弹窗不跟随布局，现场调了一下午；第四次会议是中期准备会，把所有代码合并到dev-b分支，录中期演示视频，确认了接下来一周的测试和bug修复安排。')
doc2.add_paragraph('一开始用Git分支开发，经常有人推了新代码其他人不知道，导致版本对不上。后来我们约定每次会议都先拉取最新代码，有改动及时在群里说，对接的时候两边先对齐参数再写代码，慢慢就顺畅了。我作为统筹，还要协调大家的进度，提醒谁该做什么，对接的时候两边对齐参数，虽然有点琐碎，但是也锻炼了我组织和沟通的能力。')
doc2.add_paragraph('整体来说，这次项目不仅让我把课堂上学到的拓扑排序、图论知识用到了实际项目里，更重要的是学会了怎么和别人一起协作写代码、怎么设计模块接口、怎么处理版本冲突，这些都是课堂上学不到的工程经验。虽然过程中改了很多bug，合并了很多次代码，但是最后看到整个软件能完整跑起来，从输入关系到枚举所有拓扑序到可视化高亮，那种成就感还是很强的。')

doc2.save(r'E:\tp\报告-B\个人感想-易雨杰.docx')
print('done')
