from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 封面标题
title = doc.add_heading('软件算法综合设计项目报告', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph('拓扑排序应用软件')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
info = doc.add_paragraph('学生：易雨杰  学号：2024611209  任务B：GUI主框架与交互控制')
info.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 一、需求分析
doc.add_heading('一、需求分析', level=1)
doc.add_heading('1.1 项目背景', level=2)
doc.add_paragraph('拓扑排序是对有向无环图（DAG）的所有顶点进行线性排序的算法，使得图中任意一条有向边<u,v>都满足u在v之前。本项目要求实现一个带图形界面的拓扑排序工具，用户输入"前驱→后继"形式的课程先修关系，系统自动绘制关系图，并枚举所有可能的合法拓扑排序结果。')

doc.add_heading('1.2 功能需求', level=2)
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = '功能项'
hdr[1].text = '详细要求'
funcs = [
    ('数据输入', '支持文本框直接输入、TXT文件导入，每行格式为<先修,后继>，自动过滤空行和注释'),
    ('图构建', '解析输入数据构建有向图，自动去除重复边，检测自环'),
    ('单条拓扑排序', '基于Kahn算法输出一条合法拓扑序列'),
    ('全拓扑枚举', '基于DFS回溯枚举所有合法拓扑排序，默认上限1000条'),
    ('环检测', '自动检测图中的有向环，红色高亮显示环路径'),
    ('可视化', '分层/环形两种布局，支持滚轮缩放、拖拽平移、节点拖动'),
    ('结果联动', '点击结果列表某一行，图上同步高亮该拓扑序，节点标注顺序序号'),
    ('导出', '支持排序结果导出为TXT/CSV，关系图导出为PNG'),
    ('容错', '30秒超时、计算可取消，不卡死'),
]
for f in funcs:
    row = table.add_row().cells
    row[0].text = f[0]
    row[1].text = f[1]

doc.add_heading('1.3 非功能需求', level=2)
doc.add_paragraph('- 跨平台：Java Swing实现，Windows/Linux/macOS均可运行')
doc.add_paragraph('- 零依赖：纯JDK标准库，无需安装第三方包')
doc.add_paragraph('- 易用性：操作简洁，普通用户无需培训即可使用')

# 二、概要设计
doc.add_heading('二、概要设计', level=1)
doc.add_heading('2.1 整体架构', level=2)
doc.add_paragraph('系统采用分层MVC架构，共分为四层：')
doc.add_paragraph('1. UI层：Swing组件，负责界面展示和用户交互')
doc.add_paragraph('2. 控制层：MainController，接收用户操作，调度后台线程，对接各模块')
doc.add_paragraph('3. 功能层：算法模块、可视化模块、IO模块、模型模块四个独立模块')
doc.add_paragraph('4. 模型层：图结构、顶点、边等基础数据类')

doc.add_heading('2.2 模块划分', level=2)
table2 = doc.add_table(rows=1, cols=3)
table2.style = 'Table Grid'
h = table2.rows[0].cells
h[0].text = '模块'
h[1].text = '负责人'
h[2].text = '主要职责'
mods = [
    ('A 算法模块', '骆深敏', 'Kahn排序、全拓扑枚举、环检测'),
    ('B GUI与控制模块', '易雨杰', '主窗口、控制器、结果面板、交互逻辑'),
    ('C 可视化模块', '戴燕岚', 'GraphPanel绘制、布局、缩放、导出PNG'),
    ('D IO模块', '黄佳慧', '数据解析、文件读写、导出TXT/CSV'),
    ('E 测试与交付', '吴丽梅', '功能测试、文档、会议记录、汇报材料'),
]
for m in mods:
    row = table2.add_row().cells
    row[0].text = m[0]
    row[1].text = m[1]
    row[2].text = m[2]

# 三、详细设计——B负责部分
doc.add_heading('三、详细设计（B：GUI主框架与交互控制）', level=1)
doc.add_heading('3.1 主窗口MainFrame', level=2)
doc.add_paragraph('主窗口继承JFrame，默认尺寸1200×800，采用BorderLayout布局：')
doc.add_paragraph('- 顶部：JToolBar工具栏，放置打开文件、计算、取消、导出、布局切换等按钮')
doc.add_paragraph('- 中部：JSplitPane水平分割为三栏，左栏InputPanel、中栏GraphPanel、右栏ResultPanel')
doc.add_paragraph('- 底部：StatusBar状态栏，显示节点数、边数、结果数和状态提示')

doc.add_heading('3.2 主控制器MainController', level=2)
doc.add_paragraph('MainController是整个系统的控制核心，负责：')
doc.add_paragraph('1. 接收UI事件，启动SwingWorker后台线程执行计算，避免界面卡顿')
doc.add_paragraph('2. 调用D的DataParser解析输入文本，构建Graph')
doc.add_paragraph('3. 调用A的AllTopoSorts枚举所有拓扑序，传入上限和超时控制')
doc.add_paragraph('4. 将计算结果传给C的GraphPanel画图，传给B自己的ResultPanel显示列表')
doc.add_paragraph('5. 处理结果选中事件，调用GraphPanel的setSelectedOrder高亮对应拓扑序')
doc.add_paragraph('6. 处理取消计算、导出、布局切换等操作')

doc.add_heading('3.3 结果面板ResultPanel', level=2)
doc.add_paragraph('ResultPanel负责展示所有拓扑排序结果：')
doc.add_paragraph('- 使用JList显示，分页加载，每页20条')
doc.add_paragraph('- 每条结果格式："序号. 节点1 -> 节点2 -> ... -> 节点n"')
doc.add_paragraph('- 单击选中某条结果，触发SelectionListener回调，图上同步高亮')
doc.add_paragraph('- 双击复制当前结果到剪贴板')
doc.add_paragraph('- 底部分页按钮：首页、上一页、下一页、末页、复制、清空')

doc.add_heading('3.4 交互流程', level=2)
doc.add_paragraph('用户操作流程：')
doc.add_paragraph('1. 在左栏输入或导入关系数据')
doc.add_paragraph('2. 点击"计算"按钮，后台开始解析和枚举')
doc.add_paragraph('3. 计算完成后，中栏显示关系图，右栏显示所有拓扑序')
doc.add_paragraph('4. 点击右栏某条结果，中栏图高亮该拓扑序路径，节点标注顺序')
doc.add_paragraph('5. 用户可缩放、平移、切换布局查看图，也可导出结果和图片')

doc.add_heading('3.5 B模块实现文件清单', level=2)
table3 = doc.add_table(rows=1, cols=2)
table3.style = 'Table Grid'
h3 = table3.rows[0].cells
h3[0].text = '文件路径'
h3[1].text = '功能'
files = [
    ('ui/MainFrame.java', '主窗口框架，三栏布局，工具栏'),
    ('ui/MainController.java', '主控制器，对接各模块，线程调度'),
    ('ui/ResultPanel.java', '结果列表，分页显示，选中联动'),
    ('ui/InputPanel.java', '输入面板，文本编辑，文件导入'),
    ('ui/StatusBar.java', '状态栏，统计信息，状态提示'),
    ('util/UIStyle.java', '统一UI样式，配色字体按钮'),
    ('util/ExceptionHandler.java', '统一异常弹窗处理'),
]
for f in files:
    row = table3.add_row().cells
    row[0].text = f[0]
    row[1].text = f[1]

# 四、个人感想
doc.add_heading('四、个人感想与收获', level=1)
doc.add_paragraph('作为小组的项目统筹和任务B的负责人，在这次高级算法课程项目中我收获很多。')
doc.add_paragraph('从技术层面来说，这是我第一次完整经历一个多人协作的Java项目开发流程。最开始我们五个人先一起开了几次会，确定了整体架构和接口契约，把功能拆成A到E五个模块，每个人负责一块。我负责的是GUI主框架和交互控制，一开始我以为写Swing界面很简单，就是拖一拖组件，但真正做起来才发现，要把不同人写的模块拼到一起、接口对齐、不打架，才是最花时间的。比如最开始和C对接可视化模块的时候，他写的GraphPanel和我写的MainController之间经常对不上，他改了接口我这边没更新，我这边加了功能他那边不知道，来回沟通了好几次才对齐。')
doc.add_paragraph('和其他组员的对接让我学会了怎么写清晰的接口文档，怎么提前约定好方法名、参数、返回值，而不是写完代码再凑。比如D负责导出，我们之前约定了导出的时候要传"结果是否完整"和"停止原因"两个参数，这样导出的文件才符合任务书要求，不会把截断的结果标成全部。这个过程让我理解了为什么大项目里接口设计这么重要——先把接口定死，大家各自开发，最后再拼，比边写边改效率高太多了。')
doc.add_paragraph('项目中间也遇到了不少问题，比如最开始合并代码的时候冲突一大堆，尤其是GraphPanel这个文件，C改一点我改一点，每次合并都要解决冲突。还有UI乱码的问题，一开始想用特殊符号做按钮图标，结果Swing在Windows下显示成方框，折腾了半天才发现是字体不支持，最后改成中文"重置"才解决。还有放大查看窗口不跟随布局的bug，找了半天才发现是弹新窗口的时候忘了把当前的布局模式传过去。')
doc.add_paragraph('团队合作方面，这次五个人分工明确，我们前后开了四次线上会议。第一次是项目启动会，大家一起把任务拆成五个模块，确认了每个人的分工，搭好GitHub仓库和分支；第二次会议是接口评审，把所有模块之间的方法签名、参数、返回值都定下来，冻结接口，大家开始各自开发；第三次会议是联调会，把各自写好的代码拼到一起，对接了A的算法、C的可视化、D的IO，发现了一堆接口对不上的问题，比如导出的时候缺参数、弹窗不跟随布局，现场调了一下午；第四次会议是中期准备会，把所有代码合并到dev-b分支，录中期演示视频，确认了接下来一周的测试和bug修复安排。')
doc.add_paragraph('一开始用Git分支开发，经常有人推了新代码其他人不知道，导致版本对不上。后来我们约定每次会议都先拉取最新代码，有改动及时在群里说，对接的时候两边先对齐参数再写代码，慢慢就顺畅了。我作为统筹，还要协调大家的进度，提醒谁该做什么，对接的时候两边对齐参数，虽然有点琐碎，但是也锻炼了我组织和沟通的能力。')
doc.add_paragraph('整体来说，这次项目不仅让我把课堂上学到的拓扑排序、图论知识用到了实际项目里，更重要的是学会了怎么和别人一起协作写代码、怎么设计模块接口、怎么处理版本冲突，这些都是课堂上学不到的工程经验。虽然过程中改了很多bug，合并了很多次代码，但是最后看到整个软件能完整跑起来，从输入关系到枚举所有拓扑序到可视化高亮，那种成就感还是很强的。')

doc.save(r'E:\tp\报告-B\易雨杰-任务B-最终报告部分.docx')
print('done')
