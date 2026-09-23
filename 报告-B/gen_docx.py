from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 标题
title = doc.add_heading('高级算法原理实践项目报告', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph('拓扑排序应用软件 —— 易雨杰（2024611209）任务B部分')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 一、需求分析
doc.add_heading('一、需求分析', level=1)
doc.add_heading('1.1 项目背景', level=2)
doc.add_paragraph('拓扑排序是有向无环图（DAG）中常用的线性排序算法，广泛应用于课程先修安排、任务调度、构建依赖分析等场景。本项目要求实现一个基于图形界面的拓扑排序工具，不仅能输出单条拓扑排序结果，还要枚举所有可能的合法拓扑排序，并通过可视化方式直观展示图结构和排序过程。')

doc.add_heading('1.2 功能需求', level=2)
doc.add_paragraph('根据任务书要求，结合小组讨论确定的功能需求如下：')

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '需求分类'
hdr_cells[1].text = '具体需求'
hdr_cells[2].text = '说明'

requirements = [
    ('输入功能', '文本输入', '支持直接在文本框粘贴"先修,后继"格式的关系数据'),
    ('输入功能', '文件导入', '支持从TXT文件读取关系数据，自动识别空行和注释'),
    ('解析功能', '自动建图', '解析关系数据构建有向图，自动去重、识别自环'),
    ('解析功能', '错误提示', '输入格式错误时标注行号，明确指出错误位置'),
    ('核心算法', '单条拓扑排序', '基于Kahn算法输出一条合法拓扑序列'),
    ('核心算法', '全拓扑排序枚举', '基于DFS回溯枚举所有合法拓扑排序，默认上限1000条'),
    ('核心算法', '环检测', '自动检测图中的环，有环时提示并可视化环路径'),
    ('可视化', '图渲染', '分层/环形双布局展示有向图，节点和边清晰可辨'),
    ('可视化', '交互操作', '支持滚轮缩放、拖拽平移、节点位置拖动'),
    ('可视化', '联动高亮', '点击结果列表，图上同步高亮该拓扑序的节点和顺序'),
    ('导出功能', '结果导出', '支持将排序结果导出为TXT和CSV格式'),
    ('导出功能', '图片导出', '支持将当前关系图导出为PNG图片'),
    ('容错功能', '超时保护', '计算30秒未完成自动停止，避免程序无响应'),
    ('容错功能', '取消计算', '计算过程中可随时取消，保留已生成的结果'),
]

for req in requirements:
    row_cells = table.add_row().cells
    row_cells[0].text = req[0]
    row_cells[1].text = req[1]
    row_cells[2].text = req[2]

doc.add_heading('1.3 非功能需求', level=2)
doc.add_paragraph('1. 跨平台：使用Java Swing开发，Windows/Linux/macOS均可运行')
doc.add_paragraph('2. 零依赖：纯JDK标准库实现，无需安装任何第三方依赖')
doc.add_paragraph('3. 易用性：界面简洁，操作逻辑清晰，新用户可快速上手')
doc.add_paragraph('4. 性能：100节点以内的图计算秒级完成，结果列表分页显示不卡顿')

# 二、GUI模块详细设计
doc.add_heading('二、GUI 模块详细设计', level=1)
doc.add_heading('2.1 整体架构', level=2)
doc.add_paragraph('本项目GUI采用经典的MVC架构：')
doc.add_paragraph('- 视图层（View）：Swing组件，负责界面展示和用户交互')
doc.add_paragraph('- 控制器（Controller）：MainController，接收用户操作，调用算法和IO模块，更新视图')
doc.add_paragraph('- 模型层（Model）：图结构、排序结果等数据类')
doc.add_paragraph('整体界面采用三栏式布局：顶部工具栏、中部三栏、底部状态栏。')

doc.add_heading('2.2 各组件设计', level=2)
doc.add_paragraph('主窗口 MainFrame：继承JFrame，默认尺寸1200×800，支持最大化，中部使用JSplitPane分三栏。')
doc.add_paragraph('输入面板 InputPanel：包含多行文本输入区，支持导入文件、清空、示例数据按钮。')
doc.add_paragraph('结果面板 ResultPanel：使用JList显示排序结果，每页20条，单击选中联动高亮，双击复制，底部分页按钮。')
doc.add_paragraph('关系图画布 GraphPanel（组员C负责）：支持分层/环形双布局，缩放平移，右下角悬浮缩放按钮，左上角放大查看窗口。')
doc.add_paragraph('状态栏 StatusBar：左侧显示节点数、边数、结果总数，右侧显示状态提示。')

doc.add_heading('2.3 交互流程设计', level=2)
doc.add_paragraph('1. 用户启动程序，主窗口加载完成')
doc.add_paragraph('2. 用户在输入面板粘贴数据或导入文件')
doc.add_paragraph('3. 点击"计算"按钮，后台线程解析、建图、枚举拓扑序')
doc.add_paragraph('4. 计算完成后，图和结果列表显示')
doc.add_paragraph('5. 点击结果，图上同步高亮该拓扑序')
doc.add_paragraph('6. 支持缩放、平移、切换布局查看')
doc.add_paragraph('7. 可导出结果和图片')

doc.add_heading('2.4 B模块的具体实现内容', level=2)
table2 = doc.add_table(rows=1, cols=3)
table2.style = 'Table Grid'
hdr = table2.rows[0].cells
hdr[0].text = '任务'
hdr[1].text = '实现文件'
hdr[2].text = '说明'
tasks = [
    ('主窗口框架', 'ui/MainFrame.java', '三栏布局、工具栏、状态栏'),
    ('控制器', 'ui/MainController.java', '模块对接、线程调度、事件处理'),
    ('结果面板', 'ui/ResultPanel.java', '分页显示、选中联动、复制功能'),
    ('输入面板', 'ui/InputPanel.java', '文本输入、文件导入'),
    ('状态栏', 'ui/StatusBar.java', '统计信息、状态提示'),
    ('UI样式', 'util/UIStyle.java', '统一配色、按钮样式、字体'),
    ('异常处理', 'util/ExceptionHandler.java', '统一弹窗提示错误'),
]
for t in tasks:
    r = table2.add_row().cells
    r[0].text = t[0]
    r[1].text = t[1]
    r[2].text = t[2]

# 三、个人感想
doc.add_heading('三、个人感想与收获', level=1)
doc.add_paragraph('作为小组的项目统筹和任务B的负责人，在这次高级算法课程项目中我收获很多。')
doc.add_paragraph('从技术层面来说，这是我第一次完整经历一个多人协作的Java项目开发流程。最开始我们五个人先一起开了几次会，确定了整体架构和接口契约，把功能拆成A到E五个模块，每个人负责一块。我负责的是GUI主框架和交互控制，一开始我以为写Swing界面很简单，就是拖一拖组件，但真正做起来才发现，要把不同人写的模块拼到一起、接口对齐、不打架，才是最花时间的。比如最开始和C对接可视化模块的时候，他写的GraphPanel和我写的MainController之间经常对不上，他改了接口我这边没更新，我这边加了功能他那边不知道，来回沟通了好几次才对齐。')
doc.add_paragraph('和其他组员的对接让我学会了怎么写清晰的接口文档，怎么提前约定好方法名、参数、返回值，而不是写完代码再凑。比如D负责导出，我们之前约定了导出的时候要传"结果是否完整"和"停止原因"两个参数，这样导出的文件才符合任务书要求，不会把截断的结果标成全部。这个过程让我理解了为什么大项目里接口设计这么重要——先把接口定死，大家各自开发，最后再拼，比边写边改效率高太多了。')
doc.add_paragraph('项目中间也遇到了不少问题，比如最开始合并代码的时候冲突一大堆，尤其是GraphPanel这个文件，C改一点我改一点，每次合并都要解决冲突。还有UI乱码的问题，一开始想用特殊符号做按钮图标，结果Swing在Windows下显示成方框，折腾了半天才发现是字体不支持，最后改成中文"重置"才解决。还有放大查看窗口不跟随布局的bug，找了半天才发现是弹新窗口的时候忘了把当前的布局模式传过去。')
doc.add_paragraph('团队合作方面，这次五个人分工明确，我们前后开了四次线上会议。第一次是项目启动会，大家一起把任务拆成五个模块，确认了每个人的分工，搭好GitHub仓库和分支；第二次会议是接口评审，把所有模块之间的方法签名、参数、返回值都定下来，冻结接口，大家开始各自开发；第三次会议是联调会，把各自写好的代码拼到一起，对接了A的算法、C的可视化、D的IO，发现了一堆接口对不上的问题，比如导出的时候缺参数、弹窗不跟随布局，现场调了一下午；第四次会议是中期准备会，把所有代码合并到dev-b分支，录中期演示视频，确认了接下来一周的测试和bug修复安排。')
doc.add_paragraph('一开始用Git分支开发，经常有人推了新代码其他人不知道，导致版本对不上。后来我们约定每次会议都先拉取最新代码，有改动及时在群里说，对接的时候两边先对齐参数再写代码，慢慢就顺畅了。我作为统筹，还要协调大家的进度，提醒谁该做什么，对接的时候两边对齐参数，虽然有点琐碎，但是也锻炼了我组织和沟通的能力。')
doc.add_paragraph('整体来说，这次项目不仅让我把课堂上学到的拓扑排序、图论知识用到了实际项目里，更重要的是学会了怎么和别人一起协作写代码、怎么设计模块接口、怎么处理版本冲突，这些都是课堂上学不到的工程经验。虽然过程中改了很多bug，合并了很多次代码，但是最后看到整个软件能完整跑起来，从输入关系到枚举所有拓扑序到可视化高亮，那种成就感还是很强的。')

doc.save(r'E:\tp\报告-B\易雨杰-任务B-报告部分.docx')
print('done')
