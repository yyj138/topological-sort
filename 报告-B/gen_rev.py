from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
title = doc.add_heading('修订历史记录', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = '版本号'
hdr[1].text = '日期'
hdr[2].text = '修订内容'
hdr[3].text = '修订人'

revs = [
    ('V1.0', '2026-09-17', '项目启动，完成接口契约定义、仓库搭建、各模块开发分支创建，任务分工确认', '全体'),
    ('V1.1', '2026-09-19', '完成各模块基础代码：A的核心算法、B的GUI框架、C的GraphPanel、D的IO解析，开始初步联调', '全体'),
    ('V1.2', '2026-09-20', '核心功能跑通，MVP版本完成：输入→建图→排序→可视化→结果列表全流程打通', '全体'),
    ('V1.3', '2026-09-21', '完成第一次联调：修复接口不匹配问题，对接A的算法返回结果、C的可视化高亮、D的文件导入导出', '易雨杰/骆深敏/戴燕岚/黄佳慧'),
    ('V1.4', '2026-09-22', '中期版本：完成所有核心功能，UI统一美化，修复乱码、缩放、布局切换问题，中期演示视频录制完成', '全体'),
    ('V1.5', '2026-09-22', '文档完善：README迭代更新，各模块设计文档、测试反馈文档整理完成', '易雨杰'),
]

for r in revs:
    row = table.add_row().cells
    row[0].text = r[0]
    row[1].text = r[1]
    row[2].text = r[2]
    row[3].text = r[3]

doc.save(r'E:\tp\报告-B\修订历史记录.docx')
print('done')
