from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 标题
title = doc.add_heading('拓扑排序应用工具 - B模块设计与实现', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

# 2. 需求分析
doc.add_heading('2. 需求分析', level=1)

doc.add_heading('2.1 功能需求', level=2)
doc.add_paragraph('根据任务书要求，结合小组四次会议讨论，本系统需要实现以下核心功能：')

p = doc.add_paragraph()
p.add_run('（1）关系数据输入与导入：').bold = True
p.add_run('支持文本框直接输入前驱-后继关系，格式严格遵循任务书规定的<a,b>尖括号格式，一行一个关系；支持从文本文件批量导入数据；支持表格形式编辑节点和边，方便用户调整。输入时自动过滤空行和多余空格，容错性好。')

p = doc.add_paragraph()
p.add_run('（2）关系图可视化展示：').bold = True
p.add_run('将输入的有向无环图以图形方式直观展示，节点用圆角矩形表示，边用带箭头的连线表示；支持分层布局和环形布局两种模式，方便用户根据图的大小切换查看；支持鼠标滚轮缩放、拖拽平移、节点拖动操作。')

p = doc.add_paragraph()
p.add_run('（3）拓扑排序结果枚举：').bold = True
p.add_run('计算并输出尽可能多的拓扑排序结果，默认上限1000条、30秒超时、可中途取消，三个保护机制结合，避免大结果集导致系统无响应，同时保证用户随时可以停止计算。')

p = doc.add_paragraph()
p.add_run('（4）结果分页展示与联动：').bold = True
p.add_run('以列表形式展示所有生成的拓扑排序结果，分页浏览，每页20条；点击某条结果时，关系图上对应的节点和边同步高亮，节点右上角标注顺序序号，直观展示该拓扑序的执行顺序；不同结果行用不同颜色轮换区分（绿、橙、蓝、紫、红五色循环）。')

p = doc.add_paragraph()
p.add_run('（5）环检测与高亮：').bold = True
p.add_run('自动检测输入图是否存在环，若存在则以红色高亮显示环路径，状态栏同步提示"图中存在环"，并说明环的具体路径，避免用户误以为程序出错。')

p = doc.add_paragraph()
p.add_run('（6）多格式数据导出：').bold = True
p.add_run('支持将拓扑排序结果导出为TXT或CSV文件，导出文件头需注明结果总数、是否完整、停止原因，不能把截断结果标注为全部；支持将关系图导出为PNG图片，包含完整图例，方便用户在报告和演示中使用。')

doc.add_heading('2.2 非功能需求', level=2)
p = doc.add_paragraph()
p.add_run('（1）易用性：').bold = True
p.add_run('界面布局简洁直观，常用操作都在顶部工具栏显眼位置，新手用户无需文档即可在5分钟内上手；错误提示全部用中文说明，不弹技术栈的英文堆栈信息。')

p = doc.add_paragraph()
p.add_run('（2）性能：').bold = True
p.add_run('对于15个节点以内的常规课程图，1000条拓扑排序结果的计算时间应在1秒以内；计算过程中界面不卡顿、可以正常拖动窗口和操作；界面响应时间不超过100ms。')

p = doc.add_paragraph()
p.add_run('（3）健壮性：').bold = True
p.add_run('对空输入、格式错误、环图、重复边、自环等异常情况有友好的错误提示，程序不会崩溃；计算过程中可随时取消，取消后保留已生成的部分结果，不会丢失。')

p = doc.add_paragraph()
p.add_run('（4）可扩展性：').bold = True
p.add_run('采用模块化设计，算法、GUI、可视化、IO、测试五个模块之间通过明确定义的接口通信，模块解耦，便于后续功能扩展和bug修复，改一个模块不影响其他模块。')

doc.add_heading('2.3 数据格式说明', level=2)
doc.add_paragraph('输入数据严格遵循任务书统一规定的格式：采用西文尖括号包裹的二元组，一行定义一个前驱-后继关系，例如：')
p = doc.add_paragraph()
p.add_run('<MA140, MA141>\n<MA141, CS150>\n<CS150, CS250>').font.name = 'Consolas'
doc.add_paragraph('其中尖括号内前一个元素为先修课程，后一个元素为后续课程。系统会自动解析输入文本，过滤空行和注释，构建有向图结构。测试数据包括任务书给的图1（15节点示例图）和计算机系培养方案的课程先修关系图。')

doc.add_page_break()

# 3. 详细设计
doc.add_heading('3. 详细设计', level=1)

doc.add_heading('3.2 GUI模块设计（B负责）', level=2)

doc.add_heading('3.2.1 整体架构与模块划分', level=3)
doc.add_paragraph('GUI模块采用MVC设计模式，将界面显示、业务逻辑和数据处理分离，保证代码结构清晰，便于多人协作开发。整体架构分为三层：')

p = doc.add_paragraph(style='List Bullet')
p.add_run('View层（视图组件）：').bold = True
p.add_run('MainFrame主窗口、InputPanel输入面板、ResultPanel结果面板、GraphPanel图表面板（C开发）、StatusBar状态栏、ExceptionHandler统一异常弹窗。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('Controller层（控制逻辑）：').bold = True
p.add_run('MainController主控制器，是整个系统的"中间人"，负责接收用户操作，调用算法模块和IO模块，把结果传递给视图组件更新界面。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('Model层（数据模型）：').bold = True
p.add_run('Graph图结构、EnumerationResult计算结果、ParseResult解析结果，分别由A模块和D模块提供。')

doc.add_paragraph('模块间严格通过接口通信，比如B不直接调用A的算法细节，只调用AllTopoSorts.enumerate()；也不直接调用D的文件操作细节，只调用FileManager.exportTxt()。这样每个模块内部怎么实现，别人不用关心，改起来也不会互相影响，符合面向对象的开闭原则。')

doc.add_heading('3.2.2 主窗口MainFrame', level=3)
doc.add_paragraph('主窗口采用BorderLayout布局，整体分为五个区域：')
p = doc.add_paragraph(style='List Bullet')
p.add_run('顶部工具栏：').bold = True
p.add_run('从左到右依次是打开文件、保存文件、计算、取消、导出PNG、导出结果、分层布局、环形布局、重置视图。常用操作都放在这里，用户不用翻菜单。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('左侧输入区：').bold = True
p.add_run('一个多行文本框，占左边约1/4宽度，用户直接在这里粘贴或输入关系数据，旁边有个小按钮可以打开表格编辑界面。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('中间图区：').bold = True
p.add_run('占中间大部分区域，是GraphPanel，用来显示关系图，右上角有图例，左下角有缩放指示器。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('右侧结果区：').bold = True
p.add_run('占右边约1/4宽度，是ResultPanel，列表形式展示拓扑排序结果，下面有分页按钮。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('底部状态栏：').bold = True
p.add_run('一行小字，显示当前系统状态信息，不占用太多空间。')

doc.add_heading('3.2.3 主控制器MainController', level=3)
doc.add_paragraph('MainController是GUI模块的核心，也是连接其他所有模块的枢纽，主要职责包括：')

p = doc.add_paragraph(style='List Number')
p.add_run('输入解析：').bold = True
p.add_run('拿到InputPanel的文本后，调用D模块的DataParser.parse()解析成Graph对象；如果有格式错误，直接弹中文错误提示，不进入计算流程。')

p = doc.add_paragraph(style='List Number')
p.add_run('后台任务调度：').bold = True
p.add_run('用SwingWorker在后台线程执行环检测和拓扑枚举，避免计算过程卡住界面。任务开始前记录开始时间，任务结束后计算总耗时。')

p = doc.add_paragraph(style='List Number')
p.add_run('结果分发：').bold = True
p.add_run('计算完成后，把Graph传给C模块的GraphPanel画图，把结果列表传给ResultPanel分页显示，把统计信息传给StatusBar显示。')

p = doc.add_paragraph(style='List Number')
p.add_run('用户交互处理：').bold = True
p.add_run('处理结果点击、布局切换、放大缩小、导出等操作。比如用户点某条结果行，MainController拿到结果的索引，调用GraphPanel.setSelectedOrder()在图上高亮。')

p = doc.add_paragraph(style='List Number')
p.add_run('任务状态管理：').bold = True
p.add_run('维护当前任务的引用，旧任务完成后会检查自己是不是当前任务，不是的话就不更新界面，避免连续计算时旧结果覆盖新结果。')

p = doc.add_paragraph(style='List Number')
p.add_run('状态传递：').bold = True
p.add_run('把A的StopReason枚举转换成中文文本，既显示在状态栏，也传给D的导出函数写进文件头，保证界面和导出文件的信息一致。')

doc.add_heading('3.2.4 结果面板ResultPanel', level=3)
doc.add_paragraph('结果面板是用户查看拓扑排序结果的地方，采用分页设计，每页固定20条，主要功能：')
p = doc.add_paragraph(style='List Bullet')
p.add_run('分页导航：').bold = True
p.add_run('上一页、下一页按钮，还有当前页码和总页数显示，用户可以快速翻页。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('点击联动：').bold = True
p.add_run('点击列表里某一行，立刻通知MainController，由MainController调用GraphPanel的高亮接口，图上同步显示这条拓扑序的路径和顺序。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('序号标记：').bold = True
p.add_run('每一行前面都有序号，从1开始，和图上节点的顺序徽章一一对应。')

doc.add_heading('3.2.5 状态栏StatusBar', level=3)
doc.add_paragraph('状态栏从左到右显示以下信息，让用户一眼知道当前系统状态：')
p = doc.add_paragraph(style='List Bullet')
p.add_run('节点数：').bold = True
p.add_run('当前图有多少个节点。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('边数：').bold = True
p.add_run('当前图有多少条边。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('环标记：').bold = True
p.add_run('显示"无环"或"含环"。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('序列数：').bold = True
p.add_run('这次计算一共生成了多少条拓扑排序结果。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('耗时：').bold = True
p.add_run('从点计算按钮到计算完成的总耗时，包含解析、环检测、枚举整个过程。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('停止原因：').bold = True
p.add_run('用橙色字体显示，说明为什么停止计算（达到上限/用户取消/超时/有环），正常全部完成时不显示，避免用户误以为1000条就是全部结果。')

p = doc.add_paragraph(style='List Bullet')
p.add_run('缩放比例：').bold = True
p.add_run('当前关系图的缩放百分比。')

doc.add_heading('3.2.6 核心主流程', level=3)
doc.add_paragraph('系统的完整主流程如下：')
p = doc.add_paragraph(style='List Number')
p.add_run('用户在输入框粘贴关系数据，或者点"打开文件"导入txt文件。')
p = doc.add_paragraph(style='List Number')
p.add_run('用户点击"计算"按钮，MainController开始计时，按钮变成"取消"。')
p = doc.add_paragraph(style='List Number')
p.add_run('调用D的DataParser解析输入文本，生成Graph对象；如果有格式错误，直接弹中文错误提示，流程结束。')
p = doc.add_paragraph(style='List Number')
p.add_run('调用A的CycleDetector检测环，如果有环，红色高亮环路径，状态栏提示"图中存在环"，不进行拓扑枚举。')
p = doc.add_paragraph(style='List Number')
p.add_run('在后台线程调用A的AllTopoSorts枚举所有拓扑排序，过程中支持用户点取消，达到1000条或30秒自动停。')
p = doc.add_paragraph(style='List Number')
p.add_run('计算完成后，GraphPanel画关系图，ResultPanel分页显示结果，状态栏更新所有统计信息。')
p = doc.add_paragraph(style='List Number')
p.add_run('用户点结果行，图上同步高亮对应拓扑序；用户可以点导出按钮导出结果或图片。')

doc.add_page_break()

doc.save('E:\\tp\\docs\\B负责部分-需求分析与GUI设计.docx')
print('done')
